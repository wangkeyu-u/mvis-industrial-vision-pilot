from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "AI多模态视觉算法系统_需求规格说明书_v1.0.md"
OUTPUT = ROOT / "AI多模态视觉算法系统_需求规格说明书_v1.0.docx"
ASSET_DIR = ROOT / "assets"
ARCH = ASSET_DIR / "system_architecture.png"

PAGE_WIDTH_DXA = 12240
PAGE_HEIGHT_DXA = 15840
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
ACCENT = "1F4E79"
ACCENT_2 = "2E74B5"
INK = "1F2937"
MUTED = "667085"
LIGHT = "EAF0F6"
LIGHT_GRAY = "F2F4F7"
GRID = "B8C2CC"
WHITE = "FFFFFF"
RISK = "9B1C1C"
# Named typography override: use the bundled open-source CJK face for
# deterministic DOCX-to-PDF rendering across Word and LibreOffice.
FONT_LATIN = "Noto Sans CJK SC"
FONT_CJK = "Noto Sans CJK SC"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    assert sum(widths) == CONTENT_WIDTH_DXA, (widths, sum(widths))
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    old_grid = table._tbl.tblGrid
    if old_grid is not None:
        table._tbl.remove(old_grid)
    grid = OxmlElement("w:tblGrid")
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    table._tbl.insert(1, grid)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths[idx] / 1440)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def set_run_font(run, size=None, bold=None, color=None, italic=None, mono=False) -> None:
    name = "Menlo" if mono else FONT_LATIN
    east = FONT_CJK
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), east)
    rfonts.set(qn("w:cs"), east)
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), "en-US")
    lang.set(qn("w:eastAsia"), "zh-CN")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def apply_inline(paragraph, text: str, size=None, color=None) -> None:
    pattern = re.compile(r"(\*\*.+?\*\*|`.+?`|https?://\S+)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_run_font(run, size=size, color=color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, color=color, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=(size or 11) - 0.5, color="7A3E00", mono=True)
        else:
            # Keep raw source links readable in Word without unsafe external relationship patching.
            run = paragraph.add_run(token.rstrip("。；，,)"))
            set_run_font(run, size=size, color=ACCENT_2)
            run.underline = True
            suffix = token[len(run.text):]
            if suffix:
                tail = paragraph.add_run(suffix)
                set_run_font(tail, size=size, color=color)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=size, color=color)


def add_field(paragraph, instruction: str, fallback: str = "1") -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = fallback
    fld_char3 = OxmlElement("w:fldChar")
    fld_char3.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2, text, fld_char3])
    set_run_font(run, size=9, color=MUTED)


def add_bottom_border(paragraph, color=ACCENT, size=14, space=5) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = FONT_LATIN
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.1
    pf.widow_control = True

    for style_name, size, color, before, after in (
        ("Heading 1", 16, ACCENT_2, 16, 8),
        ("Heading 2", 13, ACCENT_2, 12, 6),
        ("Heading 3", 12, ACCENT, 8, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = FONT_LATIN
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = FONT_LATIN
        style.font.size = Pt(11)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167

    code = doc.styles.add_style("Code Block", 1)
    code.font.name = "Menlo"
    code.font.size = Pt(8.5)
    code.paragraph_format.left_indent = Inches(0.18)
    code.paragraph_format.right_indent = Inches(0.18)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(6)
    code.paragraph_format.line_spacing = 1.05
    code.paragraph_format.keep_together = True
    code._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
    code._element.rPr.rFonts.set(qn("w:cs"), FONT_CJK)
    code._element.pPr.append(parse_xml(r'<w:shd {} w:fill="F5F7FA"/>'.format(nsdecls("w"))))

    quote = doc.styles.add_style("Lead Callout", 1)
    quote.font.name = FONT_LATIN
    quote.font.size = Pt(10.5)
    quote.font.bold = False
    quote.font.color.rgb = RGBColor.from_string(INK)
    quote._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
    quote.paragraph_format.left_indent = Inches(0.18)
    quote.paragraph_format.right_indent = Inches(0.18)
    quote.paragraph_format.space_before = Pt(6)
    quote.paragraph_format.space_after = Pt(8)
    quote.paragraph_format.line_spacing = 1.15
    quote._element.pPr.append(parse_xml(r'<w:shd {} w:fill="EAF0F6"/>'.format(nsdecls("w"))))


def configure_page(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def configure_header_footer(section) -> None:
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(2)
    r1 = p.add_run("MVIS-SRS-001  |  ")
    set_run_font(r1, size=8.5, color=MUTED, bold=True)
    r2 = p.add_run("轻量多模态视觉定位与结构化问答系统")
    set_run_font(r2, size=8.5, color=MUTED)
    add_bottom_border(p, color="D7DEE8", size=6, space=4)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(2)
    r = fp.add_run("内部使用  |  V1.0  |  第 ")
    set_run_font(r, size=8.5, color=MUTED)
    add_field(fp, "PAGE", "1")
    r = fp.add_run(" 页")
    set_run_font(r, size=8.5, color=MUTED)


def make_architecture() -> None:
    ASSET_DIR.mkdir(exist_ok=True)
    scale = 2
    w, h = 1500, 760
    img = Image.new("RGB", (w, h), "#FFFFFF")
    draw = ImageDraw.Draw(img)
    font_path = "/System/Library/AssetsV2/com_apple_MobileAsset_Font8/86ba2c91f017a3749571a82f2c6d890ac7ffb2fb.asset/AssetData/PingFang.ttc"
    try:
        title_font = ImageFont.truetype(font_path, 33)
        box_font = ImageFont.truetype(font_path, 26)
        small_font = ImageFont.truetype(font_path, 22)
    except OSError:
        title_font = box_font = small_font = ImageFont.load_default()

    draw.text((60, 35), "端到端逻辑架构", font=title_font, fill="#1F4E79")

    def box(x, y, bw, bh, title, detail, fill="#EAF0F6", outline="#2E74B5"):
        draw.rounded_rectangle((x, y, x + bw, y + bh), radius=18, fill=fill, outline=outline, width=3)
        tw = draw.textbbox((0, 0), title, font=box_font)[2]
        draw.text((x + (bw - tw) / 2, y + 24), title, font=box_font, fill="#163A5F")
        dw = draw.textbbox((0, 0), detail, font=small_font)[2]
        draw.text((x + (bw - dw) / 2, y + 70), detail, font=small_font, fill="#566573")

    def arrow(x1, y1, x2, y2):
        draw.line((x1, y1, x2, y2), fill="#6B7C93", width=5)
        import math
        angle = math.atan2(y2 - y1, x2 - x1)
        length = 17
        for delta in (2.55, -2.55):
            ax = x2 + length * math.cos(angle + delta)
            ay = y2 + length * math.sin(angle + delta)
            draw.line((x2, y2, ax, ay), fill="#6B7C93", width=5)

    y = 140
    bw, bh = 240, 130
    xs = [55, 345, 635, 925, 1215]
    box(xs[0], y, bw, bh, "用户 / 客户端", "图片 + 自然语言")
    box(xs[1], y, bw, bh, "输入与任务解析", "校验 · 预处理 · 路由")
    box(xs[2], y, bw, bh, "主 VLM 推理", "Qwen3-VL + LoRA", fill="#DDEBF7")
    box(xs[3], y, bw, bh, "融合与策略", "冲突 · 置信度 · 拒答")
    box(xs[4], y, bw, bh, "验证与输出", "Schema · API · 可视化")
    for i in range(4):
        arrow(xs[i] + bw, y + bh // 2, xs[i + 1] - 10, y + bh // 2)

    box(635, 350, 530, 125, "专用视觉适配器", "RF-DETR Nano / Florence-2-base", fill="#F8F4E8", outline="#B7791F")
    arrow(800, 350, 770, 275)
    arrow(1010, 350, 1045, 275)

    draw.rounded_rectangle((55, 560, 1455, 700), radius=18, fill="#F4F6F9", outline="#AAB7C4", width=3)
    supports = [
        (105, "数据与版本", "manifest · 许可 · 划分"),
        (510, "实验与评测", "基线 · 消融 · 切片"),
        (915, "模型与产物", "权重 · 配置 · 模型卡"),
        (1210, "日志与监控", "延迟 · 内存 · 错误"),
    ]
    for x, title, detail in supports:
        draw.text((x, 590), title, font=box_font, fill="#344054")
        draw.text((x, 638), detail, font=small_font, fill="#667085")

    img.resize((w // scale, h // scale), Image.Resampling.LANCZOS).save(ARCH, quality=95)


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("需求规格说明书")
    set_run_font(r, size=11, color=ACCENT_2, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("轻量多模态视觉定位与\n结构化问答系统")
    set_run_font(r, size=25, color=ACCENT, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(28)
    r = p.add_run("企业级需求规格说明书（SRS / PRD）")
    set_run_font(r, size=14, color=MUTED)
    add_bottom_border(p, color=ACCENT_2, size=16, space=8)

    meta = [
        ("文档编号", "MVIS-SRS-001"),
        ("版本 / 状态", "V1.0 / 需求基线草案"),
        ("发布日期", "2026-08-08"),
        ("适用阶段", "作品集项目 / MVP / 算法面试"),
        ("文档负责人", "项目负责人（待指定）"),
        ("密级", "内部使用"),
    ]
    table = doc.add_table(rows=len(meta), cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [1800, 7560])
    for row, (label, value) in zip(table.rows, meta):
        set_cell_shading(row.cells[0], LIGHT)
        pr = row.cells[0].paragraphs[0]
        pr.paragraph_format.space_after = Pt(0)
        rr = pr.add_run(label)
        set_run_font(rr, size=10, color=ACCENT, bold=True)
        pv = row.cells[1].paragraphs[0]
        pv.paragraph_format.space_after = Pt(0)
        rv = pv.add_run(value)
        set_run_font(rv, size=10.5, color=INK)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(82)
    p.paragraph_format.space_after = Pt(4)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("文档用途")
    set_run_font(r, size=9, color=MUTED, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("用于统一产品、算法、数据、工程、评测与验收口径；作为项目实施和技术面试展示的需求基线。")
    set_run_font(r, size=10, color=MUTED)

    doc.add_page_break()


def parse_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
            continue
        rows.append(cells)
    return rows


def choose_widths(headers: list[str]) -> list[int]:
    n = len(headers)
    joined = "|".join(headers)
    if n == 2:
        return [1800, 7560]
    if n == 3:
        if "HTTP" in headers:
            return [3450, 750, 5160]
        if any(k in joined for k in ("类型", "必填", "优先级")):
            return [1900, 1700, 5760]
        return [1700, 3830, 3830]
    if n == 4:
        if headers[:3] == ["字段", "类型", "必填"]:
            return [1600, 1700, 900, 5160]
        if headers[0] in ("ID", "编号") or "优先级" in headers:
            return [1250, 700, 3950, 3460]
        if "角色" in headers and "权限" in joined:
            return [1600, 4200, 3560]
        return [1450, 1750, 3300, 2860]
    if n == 5:
        if headers[0] == "活动":
            return [1100, 2050, 2050, 2050, 2110]
        if "决策项" in headers:
            return [750, 1750, 2650, 1800, 2410]
        return [750, 1600, 1100, 3300, 2610]
    base = CONTENT_WIDTH_DXA // n
    widths = [base] * n
    widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
    return widths


def create_decimal_num_id(doc: Document) -> int:
    """Create an independent, real Word decimal list that starts at 1."""
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [
        int(node.get(qn("w:numId")))
        for node in numbering.findall(qn("w:num"))
    ]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    for tag, value in (("w:start", "1"), ("w:numFmt", "decimal"), ("w:lvlText", "%1."), ("w:lvlJc", "left")):
        element = OxmlElement(tag)
        element.set(qn("w:val"), value)
        level.append(element)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    p_pr.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "720")
    indent.set(qn("w:hanging"), "360")
    p_pr.append(indent)
    level.append(p_pr)
    abstract.append(level)

    first_num = numbering.find(qn("w:num"))
    if first_num is None:
        numbering.append(abstract)
    else:
        numbering.insert(numbering.index(first_num), abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    old = p_pr.find(qn("w:numPr"))
    if old is not None:
        p_pr.remove(old)
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_ref = OxmlElement("w:numId")
    num_ref.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_ref)
    p_pr.append(num_pr)


def add_markdown_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = len(rows[0])
    rows = [r[:cols] + [""] * max(0, cols - len(r)) for r in rows]
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    set_table_geometry(table, choose_widths(rows[0]))
    set_repeat_table_header(table.rows[0])
    for r_idx, (word_row, values) in enumerate(zip(table.rows, rows)):
        for c_idx, (cell, value) in enumerate(zip(word_row.cells, values)):
            if r_idx == 0:
                set_cell_shading(cell, LIGHT_GRAY)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.06
            if c_idx in (0, 1) and len(value) < 18:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cell_size = 8.7 if cols >= 4 else 9.2
            if "HTTP" in rows[0] and c_idx == 0 and len(value) > 20:
                cell_size = 8.0
            apply_inline(p, value, size=cell_size, color=INK)
            for run in p.runs:
                if r_idx == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(ACCENT)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.space_before = Pt(0)


def add_image(doc: Document, alt: str, path: Path) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run()
    inline = run.add_picture(str(path), width=Inches(6.25))
    doc_pr = inline._inline.docPr
    doc_pr.set("descr", alt)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    rr = cap.add_run(f"图 1  {alt}")
    set_run_font(rr, size=9, color=MUTED, italic=True)


def build_body(doc: Document, markdown: str) -> None:
    lines = markdown.splitlines()
    # Everything before the first explicit page-break marker is represented by the custom cover.
    start = next(i for i, line in enumerate(lines) if line.strip() == "<!-- PAGEBREAK -->") + 1
    lines = lines[start:]
    i = 0
    in_code = False
    code_lines: list[str] = []
    active_num_id: int | None = None
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()
        is_number_line = bool(re.match(r"^\d+\.\s+", stripped))
        if not is_number_line:
            active_num_id = None

        if stripped.startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                p = doc.add_paragraph(style="Code Block")
                for idx, code_line in enumerate(code_lines):
                    if idx:
                        p.add_run("\n")
                    rr = p.add_run(code_line)
                    set_run_font(rr, size=8.5, color=INK, mono=True)
                in_code = False
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if not stripped:
            i += 1
            continue
        if stripped == "<!-- PAGEBREAK -->":
            doc.add_page_break()
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-{3,}", lines[i + 1]):
            table_lines = [line]
            i += 1
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_markdown_table(doc, parse_table(table_lines))
            continue

        image_match = re.fullmatch(r"!\[(.+?)\]\((.+?)\)", stripped)
        if image_match:
            add_image(doc, image_match.group(1), ROOT / image_match.group(2))
            i += 1
            continue

        head = re.match(r"^(#{2,4})\s+(.+)$", stripped)
        if head:
            level = len(head.group(1)) - 1
            p = doc.add_paragraph(style=f"Heading {min(level, 3)}")
            apply_inline(p, head.group(2), color=ACCENT_2 if level < 3 else ACCENT)
            i += 1
            continue

        if stripped.startswith("> "):
            p = doc.add_paragraph(style="Lead Callout")
            apply_inline(p, stripped[2:], size=10.5, color=INK)
            i += 1
            continue

        if re.match(r"^-\s+", stripped):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.keep_together = True
            apply_inline(p, re.sub(r"^-\s+", "", stripped))
            i += 1
            continue

        if re.match(r"^\d+\.\s+", stripped):
            if active_num_id is None:
                active_num_id = create_decimal_num_id(doc)
            p = doc.add_paragraph(style="List Number")
            apply_numbering(p, active_num_id)
            p.paragraph_format.keep_together = True
            apply_inline(p, re.sub(r"^\d+\.\s+", "", stripped))
            i += 1
            continue

        # Join wrapped prose lines until the next structural marker.
        paragraph_lines = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if not nxt:
                i += 1
                break
            if (
                nxt.startswith(("#", "|", "- ", "> ", "```", "![", "<!--"))
                or re.match(r"^\d+\.\s+", nxt)
            ):
                break
            paragraph_lines.append(nxt)
            i += 1
        p = doc.add_paragraph()
        apply_inline(p, " ".join(paragraph_lines))


def audit_tables(doc: Document) -> None:
    for table in doc.tables:
        tbl_pr = table._tbl.tblPr
        assert tbl_pr.find(qn("w:tblW")) is not None
        assert tbl_pr.find(qn("w:tblInd")) is not None
        assert table._tbl.tblGrid is not None
        for row in table.rows:
            for cell in row.cells:
                assert cell._tc.get_or_add_tcPr().find(qn("w:tcW")) is not None


def main() -> None:
    make_architecture()
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    configure_page(section)
    configure_header_footer(section)
    add_cover(doc)
    build_body(doc, SOURCE.read_text(encoding="utf-8"))

    props = doc.core_properties
    props.title = "轻量多模态视觉定位与结构化问答系统 - 企业级需求规格说明书"
    props.subject = "多模态 AI 算法项目需求基线"
    props.author = "项目组"
    props.keywords = "VLM,Qwen3-VL,RF-DETR,LoRA,QLoRA,需求规格说明书"
    props.comments = "MVIS-SRS-001 V1.0"

    settings = doc.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")

    audit_tables(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
